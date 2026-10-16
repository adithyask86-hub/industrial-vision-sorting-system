from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).parent / 'docs' / 'Review2_Code_Report.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(9.5)
styles['Title'].font.name = 'Aptos Display'; styles['Title'].font.size = Pt(22); styles['Title'].font.bold = True
for name in ['Heading 1','Heading 2']:
    styles[name].font.name = 'Aptos Display'
    styles[name].font.color.rgb = RGBColor(20, 55, 85)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('REVIEW 2 — CODE DEVELOPMENT REPORT'); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor(15,45,70)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Real-Time Vision-Based Industrial Quality Inspection and Shape Sorting System'); r.bold=True; r.font.size=Pt(13)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Course: Microprocessor, Microcontroller and Interfacing Techniques | Course Code: 21ECC301P\n')
p.add_run('Team: Adithya Sajeev Kumar (RA2411004010326), Abhinandh M (RA2411004010336), Dharun Prakash (RA2411004010342)')

doc.add_heading('1. Project objective', level=1)
doc.add_paragraph('The project develops a low-cost prototype that inspects a single object in a controlled camera region, classifies its silhouette as a circle, triangle, square/rectangle, or unknown, applies configurable geometric defect rules, and sends a compact decision code to a microcontroller for sorting or rejection. The implementation follows the Review 1 architecture: the laptop/Raspberry Pi performs image processing while the MCU handles motor and actuator interfacing.')

doc.add_heading('2. Software architecture', level=1)
t = doc.add_table(rows=1, cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Light Shading Accent 1'
for c,txt in zip(t.rows[0].cells,['Module','Implementation','Responsibility']): c.text=txt
rows=[('Vision application','Python 3, OpenCV, NumPy','Camera acquisition, ROI, filtering, thresholding, morphology, contours, feature extraction, shape classification, rule-based inspection, annotated preview.'),('Communication','USB serial / pySerial','Sends C, S, T, or D as a newline-terminated ASCII command at 115200 baud.'),('MCU firmware','Arduino-compatible C++ / Servo library','Controls motor-driver signals, reject/sort servos, LEDs, buzzer, and good/defect counters.'),('Configuration and tests','JSON + unittest','Allows threshold calibration and checks core shape-detection behavior on synthetic images.')]
for row in rows:
    cells=t.add_row().cells
    for c,txt in zip(cells,row): c.text=txt
    for c in cells: c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

doc.add_heading('3. Image-processing algorithm', level=1)
steps=[
('Image acquisition','OpenCV reads frames from a USB webcam or compatible camera device.'),
('Region of interest (ROI)','A normalized rectangle restricts processing to the inspection zone.'),
('Noise reduction','A 5×5 Gaussian blur reduces high-frequency noise.'),
('Segmentation','Otsu thresholding creates two candidate polarities; the application ignores contours touching the ROI boundary to avoid treating the background as the object.'),
('Morphology','Opening suppresses small specks and closing fills small gaps in the segmented foreground.'),
('Contour and features','The largest valid contour is measured for area, perimeter, circularity, convex-hull solidity, bounding-box aspect ratio, polygon vertices, and nested holes.'),
('Decision rules','The object is flagged for configured size violations, unknown/deformed outline, low solidity, unexpected holes, or circle circularity below threshold.'),
('Output','The result is drawn on the preview and mapped to a serial code for the MCU.')]
for title,body in steps:
    p=doc.add_paragraph(style='List Number'); p.add_run(title+': ').bold=True; p.add_run(body)

doc.add_heading('4. Defect rules and important interpretation', level=1)
doc.add_paragraph('The baseline detects geometric/silhouette anomalies under controlled lighting. It is not a universal industrial defect detector. A crack that does not change the outer silhouette may not be detected. Surface scratches, subtle material defects, overlapping objects, and changing illumination require additional methods and validation. CNN/ML is intentionally not included because Review 1 identifies it as a future enhancement.')

t=doc.add_table(rows=1, cols=3); t.style='Light Shading Accent 1'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for c,txt in zip(t.rows[0].cells,['Feature / rule','Purpose','Calibration note']): c.text=txt
for row in [
('Area min/max','Reject objects outside expected pixel-size range.','Depends on camera resolution, distance, and object dimensions.'),
('Solidity','Flags strongly concave or missing-portion silhouettes.','Tune against examples of acceptable and defective parts.'),
('Hole count','Flags enclosed holes when not allowed.','Small holes/noise are filtered by a relative area threshold.'),
('Polygon vertices + circularity','Classifies common silhouettes and flags unknown/deformed outlines.','Works best with matte objects and a contrasting background.')]:
    cells=t.add_row().cells
    for c,txt in zip(cells,row): c.text=txt

doc.add_heading('5. Serial protocol', level=1)
t=doc.add_table(rows=1, cols=3); t.style='Light Shading Accent 1'
for c,txt in zip(t.rows[0].cells,['Code','Vision meaning','MCU action']): c.text=txt
for row in [('C','Good circle','Move sorting servo to circle-bin angle'),('S','Good square/rectangle','Move sorting servo to square-bin angle'),('T','Good triangle','Move sorting servo to triangle-bin angle'),('D','Defective or unknown','Move rejection servo')]:
    cells=t.add_row().cells
    for c,txt in zip(cells,row): c.text=txt

doc.add_paragraph('The MCU returns ACK messages with good/defect counters or ERR messages. The current baseline uses a simple serial result protocol. Although the IR input pin is declared in firmware, it is not yet used as a trigger/handshake. The Python application sends when the visible result signature changes; for a moving conveyor, production-grade object tracking requires a sensor-triggered request/reply sequence and position timing/encoder support.')

doc.add_heading('6. Files delivered', level=1)
for x in [
'vision/inspector.py — OpenCV processing, live preview, still-image mode, and optional serial transmission.',
'config.json — camera size, ROI, area limits, morphology, solidity, hole, and circularity thresholds.',
'firmware/vision_sorter.ino — MCU motor/servo/indicator control and serial parser.',
'requirements.txt — Python dependencies.',
'tests/test_inspector.py — four synthetic-image smoke tests.',
'README.md — setup, wiring cautions, commands, calibration, and limitations.']:
    doc.add_paragraph(x, style='List Bullet')

doc.add_heading('7. Verification performed', level=1)
doc.add_paragraph('The Python smoke-test suite was executed in the development environment. All four tests passed: circle classification, square/rectangle classification, triangle classification, and empty-frame handling. These are algorithm-level synthetic tests only; they do not establish real-camera accuracy, physical sorting accuracy, processing-time guarantees, or successful hardware actuation. Those must be measured on the actual prototype.')

doc.add_heading('8. Setup and execution', level=1)
for line in [
'Create and activate a Python virtual environment, then run: pip install -r requirements.txt',
'Live camera only: python vision/inspector.py --config config.json --camera 0',
'With MCU: python vision/inspector.py --config config.json --camera 0 --serial COM4 (replace COM4 with the correct port)',
'Still image: python vision/inspector.py --config config.json --snapshot path/to/image.jpg',
'Tests: python -m unittest discover -s tests -v',
'Open firmware/vision_sorter.ino in Arduino IDE, select the correct board/port, and upload.']:
    doc.add_paragraph(line, style='List Bullet')

doc.add_heading('9. Safety and hardware notes', level=1)
doc.add_paragraph('Do not connect motors or servo power directly to MCU pins or rely on the USB port for motor loads. Use a motor driver rated for the motor stall current and a suitable external regulated supply for servos/motors. Tie the external supply ground to MCU ground. Verify pin mappings, servo angles, belt direction, emergency stop, and mechanical clearances before running the conveyor. Start with the belt disconnected and test actuators individually.')

doc.add_heading('10. Remaining work before final review', level=1)
for x in [
'Calibrate ROI and thresholds using the actual camera, lighting, and sample objects.',
'Collect labeled good/defective samples and calculate defect precision/recall, shape accuracy, sorting accuracy, and per-object processing time.',
'Integrate IR-triggered inspection and object-to-actuator timing; verify one decision per physical object.',
'Adapt the firmware to the exact selected MCU if STM32 rather than Arduino-compatible hardware is used.',
'Add LCD/OLED support if required by the final prototype; the current baseline reports counters over serial.',
'Capture demo evidence: live annotated frames, serial logs, actuator video, test table, and measured results.']:
    doc.add_paragraph(x, style='List Bullet')

p=doc.add_paragraph(); p.add_run('Status statement for Review 2: ').bold=True
p.add_run('The code baseline covers the main vision pipeline, shape classification, configurable geometric inspection rules, serial decision output, and MCU actuation logic. Hardware integration, calibration, real-object validation, and performance measurement remain to be demonstrated. Do not describe the system as 75% complete unless the team can substantiate that percentage against its own task checklist.')

# Footer
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('21ECC301P | Review 2 Code Development Report')

doc.save(OUT)
print(OUT)
